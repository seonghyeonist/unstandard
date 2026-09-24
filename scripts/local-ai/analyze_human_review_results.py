"""Merge local blinded review workbooks and emit aggregate agreement metrics.

This tool reads only Pair ID and Reviewer label from each reviewer packet. It
never copies question text, answer text, notes, or source row IDs to its output
or stdout. Keep its inputs and outputs on the operator's local machine.
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

LABELS = ("PASS", "REVIEW", "REJECT")
INPUT_FILES = {
    "A": "reviewer_A_225.xlsx",
    "B": "reviewer_B_225.xlsx",
    "crosswalk": "controller_crosswalk_225.xlsx",
}


def _read_reviews(path: Path) -> dict[str, str]:
    book = load_workbook(path, read_only=False, data_only=True)
    try:
        if book.sheetnames != ["Review", "Instructions"] or any(
            sheet.sheet_state != "visible" for sheet in book.worksheets
        ):
            raise ValueError("review workbook contains an unexpected or hidden sheet")
        sheet = book["Review"]
        headers = [sheet.cell(1, col).value for col in range(1, sheet.max_column + 1)]
        if headers != ["Pair ID", "Question text", "Answer text", "Reviewer label", "Reviewer notes"]:
            raise ValueError("review workbook headers do not match the blinded packet")
        if sheet.max_column != 5 or any(dim.hidden for dim in sheet.column_dimensions.values()):
            raise ValueError("review workbook contains unexpected or hidden columns")
        results: dict[str, str] = {}
        for row in range(2, sheet.max_row + 1):
            pair_id = sheet.cell(row, 1).value
            label = sheet.cell(row, 4).value
            if not isinstance(pair_id, str) or not pair_id:
                raise ValueError("review workbook contains an invalid Pair ID")
            if pair_id in results:
                raise ValueError("review workbook contains duplicate Pair IDs")
            if label not in LABELS:
                raise ValueError("review workbook is incomplete or contains an invalid label")
            results[pair_id] = label
        return results
    finally:
        book.close()


def _read_categories(path: Path) -> dict[str, str]:
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Controller crosswalk" not in book.sheetnames:
            raise ValueError("controller workbook is missing its crosswalk sheet")
        sheet = book["Controller crosswalk"]
        headers = [sheet.cell(1, col).value for col in range(1, sheet.max_column + 1)]
        if headers != [
            "Pair ID", "Source row IDs", "Category", "Question level",
            "Recommended label", "Recommended path", "Expected Depth Score range",
        ]:
            raise ValueError("controller crosswalk headers do not match the local packet")
        id_col, category_col = headers.index("Pair ID") + 1, headers.index("Category") + 1
        categories: dict[str, str] = {}
        for row in range(2, sheet.max_row + 1):
            pair_id, category = sheet.cell(row, id_col).value, sheet.cell(row, category_col).value
            if not isinstance(pair_id, str) or not pair_id or not isinstance(category, str) or not category:
                raise ValueError("controller crosswalk contains an invalid entry")
            if pair_id in categories:
                raise ValueError("controller crosswalk contains duplicate Pair IDs")
            categories[pair_id] = category
        return categories
    finally:
        book.close()


def _confusion(a: list[str], b: list[str]) -> dict[str, dict[str, int]]:
    return {left: {right: sum(x == left and y == right for x, y in zip(a, b)) for right in LABELS} for left in LABELS}


def _metrics(a: list[str], b: list[str]) -> dict[str, Any]:
    n = len(a)
    if n == 0 or len(b) != n:
        raise ValueError("cannot score an empty or mismatched review group")
    matrix = _confusion(a, b)
    agreements = sum(matrix[label][label] for label in LABELS)
    counts_a, counts_b = Counter(a), Counter(b)
    observed = agreements / n
    expected = sum(counts_a[label] * counts_b[label] for label in LABELS) / (n * n)
    kappa = None if expected == 1 else (observed - expected) / (1 - expected)
    return {
        "n": n,
        "agreements": agreements,
        "disagreements": n - agreements,
        "raw_agreement": observed,
        "expected_agreement": expected,
        "cohen_kappa": kappa,
        "reviewer_a_counts": {label: counts_a[label] for label in LABELS},
        "reviewer_b_counts": {label: counts_b[label] for label in LABELS},
        "confusion_matrix_reviewer_a_rows_reviewer_b_columns": matrix,
    }


def analyze(reviewer_a: dict[str, str], reviewer_b: dict[str, str], categories: dict[str, str]) -> tuple[dict[str, Any], list[dict[str, str]]]:
    ids_a, ids_b, ids_c = set(reviewer_a), set(reviewer_b), set(categories)
    if ids_a != ids_b or ids_a != ids_c:
        raise ValueError("reviewer packets and controller crosswalk have different Pair ID sets")
    if len(ids_a) < 200:
        raise ValueError("at least 200 unique pairs are required for a complete human-review result")
    ordered_ids = list(reviewer_a)
    a_labels = [reviewer_a[pair_id] for pair_id in ordered_ids]
    b_labels = [reviewer_b[pair_id] for pair_id in ordered_ids]
    stats: dict[str, Any] = {"overall": _metrics(a_labels, b_labels), "by_category": {}}
    for category in sorted(set(categories.values())):
        positions = [i for i, pair_id in enumerate(ordered_ids) if categories[pair_id] == category]
        stats["by_category"][category] = _metrics(
            [a_labels[i] for i in positions], [b_labels[i] for i in positions]
        )
    merged = [
        {
            "pair_id": pair_id,
            "category": categories[pair_id],
            "reviewer_a_label": reviewer_a[pair_id],
            "reviewer_b_label": reviewer_b[pair_id],
            "final_label": reviewer_a[pair_id] if reviewer_a[pair_id] == reviewer_b[pair_id] else "",
            "adjudication_required": "NO" if reviewer_a[pair_id] == reviewer_b[pair_id] else "YES",
        }
        for pair_id in ordered_ids
    ]
    stats["final_labels_before_adjudication"] = dict(Counter(
        row["final_label"] for row in merged if row["final_label"]
    ))
    stats["adjudication_required"] = sum(row["adjudication_required"] == "YES" for row in merged)
    return stats, merged


def _write_outputs(directory: Path, stats: dict[str, Any], merged: list[dict[str, str]]) -> tuple[Path, Path]:
    metrics_path = directory / "human_review_metrics.json"
    merged_path = directory / "human_review_merged.xlsx"
    if metrics_path.exists() or merged_path.exists():
        raise ValueError("result outputs already exist; choose a new local output directory")
    metrics_path.write_text(json.dumps(stats, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    book = Workbook()
    sheet = book.active
    sheet.title = "Merged results"
    headers = ["Pair ID", "Category", "Reviewer A label", "Reviewer B label", "Final label", "Adjudication required"]
    sheet.append(headers)
    fill = PatternFill("solid", fgColor="17365D")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
    for row in merged:
        sheet.append([
            row["pair_id"], row["category"], row["reviewer_a_label"], row["reviewer_b_label"],
            row["final_label"], row["adjudication_required"],
        ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:F{sheet.max_row}"
    book.save(merged_path)
    os.chmod(metrics_path, 0o600)
    os.chmod(merged_path, 0o600)
    return metrics_path, merged_path


def process(directory: Path) -> dict[str, Any]:
    reviews_a = _read_reviews(directory / INPUT_FILES["A"])
    reviews_b = _read_reviews(directory / INPUT_FILES["B"])
    categories = _read_categories(directory / INPUT_FILES["crosswalk"])
    stats, merged = analyze(reviews_a, reviews_b, categories)
    metrics_path, merged_path = _write_outputs(directory, stats, merged)
    return {
        "status": "COMPLETE",
        "n": stats["overall"]["n"],
        "agreements": stats["overall"]["agreements"],
        "disagreements": stats["overall"]["disagreements"],
        "raw_agreement": stats["overall"]["raw_agreement"],
        "cohen_kappa": stats["overall"]["cohen_kappa"],
        "adjudication_required": stats["adjudication_required"],
        "outputs_written": [metrics_path.name, merged_path.name],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--directory",
        type=Path,
        default=os.environ.get("UNSTANDARD_LABELING_PACKET_DIR"),
        help="operator-local packet directory (or UNSTANDARD_LABELING_PACKET_DIR)",
    )
    args = parser.parse_args()
    if not args.directory:
        print(json.dumps({"status": "BLOCKED", "reason": "local packet directory is required"}))
        return 2
    try:
        result = process(args.directory)
    except (OSError, ValueError, KeyError) as error:
        # Error messages are deliberately generic and never include row content.
        print(json.dumps({"status": "BLOCKED", "reason": str(error)}, ensure_ascii=False))
        return 2
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
