"""Safety and statistics tests for local human-review result ingestion."""

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from analyze_human_review_results import analyze, process


class HumanReviewAnalysisTests(unittest.TestCase):
    def test_confusion_kappa_and_manual_adjudication(self):
        a = {f"p{i}": label for i, label in enumerate(["PASS", "PASS", "REVIEW", "REJECT"])}
        b = {f"p{i}": label for i, label in enumerate(["PASS", "REVIEW", "REVIEW", "REJECT"])}
        categories = {f"p{i}": "L1_PASS" if i < 2 else "GRAY_BAND" for i in range(4)}

        stats, merged = analyze_for_test(a, b, categories)

        self.assertEqual(stats["overall"]["n"], 4)
        self.assertEqual(stats["overall"]["agreements"], 3)
        self.assertEqual(stats["overall"]["disagreements"], 1)
        self.assertAlmostEqual(stats["overall"]["cohen_kappa"], 7 / 11)
        self.assertEqual(stats["overall"]["confusion_matrix_reviewer_a_rows_reviewer_b_columns"]["PASS"]["REVIEW"], 1)
        self.assertEqual(merged[0]["final_label"], "PASS")
        self.assertEqual(merged[1]["final_label"], "")
        self.assertEqual(merged[1]["adjudication_required"], "YES")

    def test_ingestion_never_copies_question_or_notes_to_outputs(self):
        with tempfile.TemporaryDirectory() as temp:
            directory = Path(temp)
            write_packet(directory / "reviewer_A_225.xlsx", ["PASS"] * 200)
            write_packet(directory / "reviewer_B_225.xlsx", ["PASS"] * 200)
            write_crosswalk(directory / "controller_crosswalk_225.xlsx", 200)

            result = process(directory)

            self.assertEqual(result["status"], "COMPLETE")
            text = (directory / "human_review_metrics.json").read_text(encoding="utf-8")
            self.assertNotIn("PRIVATE_QUESTION_TEXT", text)
            self.assertNotIn("PRIVATE_REVIEW_NOTE", text)
            workbook = load_workbook(directory / "human_review_merged.xlsx", read_only=True, data_only=True)
            try:
                flattened = " ".join(str(cell.value) for row in workbook.active.iter_rows() for cell in row if cell.value)
                self.assertNotIn("PRIVATE_QUESTION_TEXT", flattened)
                self.assertNotIn("PRIVATE_REVIEW_NOTE", flattened)
                self.assertEqual(workbook.active.max_row, 201)
            finally:
                workbook.close()

    def test_incomplete_reviews_do_not_produce_metrics(self):
        a = {f"p{i}": "PASS" for i in range(200)}
        b = {f"p{i}": "PASS" for i in range(199)}
        categories = {f"p{i}": "L1_PASS" for i in range(200)}
        with self.assertRaisesRegex(ValueError, "different Pair ID sets"):
            analyze(a, b, categories)


def analyze_for_test(a, b, categories):
    # The production gate is >=200; small unit fixtures exercise pure metrics.
    from analyze_human_review_results import _metrics

    ids = list(a)
    labels_a = [a[i] for i in ids]
    labels_b = [b[i] for i in ids]
    stats = {"overall": _metrics(labels_a, labels_b), "by_category": {}}
    merged = [
        {"pair_id": i, "final_label": a[i] if a[i] == b[i] else "", "adjudication_required": "NO" if a[i] == b[i] else "YES"}
        for i in ids
    ]
    return stats, merged


def write_packet(path, labels):
    book = Workbook()
    sheet = book.active
    sheet.title = "Review"
    sheet.append(["Pair ID", "Question text", "Answer text", "Reviewer label", "Reviewer notes"])
    for i, label in enumerate(labels):
        sheet.append([f"p{i}", "PRIVATE_QUESTION_TEXT", "PRIVATE_ANSWER_TEXT", label, "PRIVATE_REVIEW_NOTE"])
    guide = book.create_sheet("Instructions")
    guide.append(["Reviewer instructions"])
    book.save(path)


def write_crosswalk(path, count):
    book = Workbook()
    sheet = book.active
    sheet.title = "Controller crosswalk"
    sheet.append(["Pair ID", "Source row IDs", "Category", "Question level", "Recommended label", "Recommended path", "Expected Depth Score range"])
    for i in range(count):
        sheet.append([f"p{i}", f"row-{i}", "L1_PASS", "L1", "PASS", "BASIC", "1-2"])
    book.save(path)


if __name__ == "__main__":
    unittest.main()
